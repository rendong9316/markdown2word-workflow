import copy
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============ 可调参数：三列宽度（厘米）============
# 顺序：新空列、原来的第1列(公式)、原来的第2列
COL_WIDTHS_CM = [1.5, 13, 1.5]


def cm_to_twips(cm: float) -> int:
    # 1 inch = 1440 twips, 1 cm ≈ 567 twips
    return int(round(cm * 567))


def set_gridcol_width(gridCol, twips: int):
    gridCol.set(qn('w:w'), str(twips))


def set_cell_width(tc, twips: int):
    """给 <w:tc> 写 <w:tcPr><w:tcW w:w=.. w:type='dxa'/>"""
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(twips))
    tcW.set(qn('w:type'), 'dxa')


def set_fixed_layout(tblEl):
    """把表格布局改为 fixed，列宽才会被尊重"""
    tblPr = tblEl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tblEl.insert(0, tblPr)
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')


doc = Document("输出15.docx")

widths_twips = [cm_to_twips(c) for c in COL_WIDTHS_CM]

for tbl in doc.tables:
    if len(tbl.columns) != 2:
        continue

    # 粗略判断是否是"公式表"：任一单元格里含 OMML 公式
    has_math = any(
        cell._tc.findall('.//' + qn('m:oMath'))
        for row in tbl.rows for cell in row.cells
    )
    if not has_math:
        continue

    tblEl = tbl._tbl
    grid = tblEl.find(qn('w:tblGrid'))

    # 1) 最左侧插入一个新的 gridCol
    newGridCol = copy.deepcopy(grid[0])
    grid.insert(0, newGridCol)

    # 2) 重新分配三列宽度（覆盖原来的 w:w）
    gridCols = grid.findall(qn('w:gridCol'))
    for gc, w in zip(gridCols, widths_twips):
        set_gridcol_width(gc, w)

    # 3) 每行最前面插入空 tc，并给整行三个单元格重设宽度
    for tr in tblEl.findall(qn('w:tr')):
        tcs = tr.findall(qn('w:tc'))
        newTc = copy.deepcopy(tcs[0])
        # 清空除 tcPr 外的内容
        for child in list(newTc):
            if child.tag != qn('w:tcPr'):
                newTc.remove(child)
        newTc.append(OxmlElement('w:p'))   # 占位空段落
        tr.insert(0, newTc)

        # 重设本行三个单元格宽度
        tcs = tr.findall(qn('w:tc'))       # 重新取，现在是 3 个
        for tc, w in zip(tcs, widths_twips):
            set_cell_width(tc, w)

    # 4) 公式列（现在的第 2 列，索引 1）内容居中
    for row in tbl.rows:
        if len(row.cells) >= 2:
            for p in row.cells[1].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 5) 固定布局 + 表格整体居中
    set_fixed_layout(tblEl)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

doc.save("输出3_fixed.docx")